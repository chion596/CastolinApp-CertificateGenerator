import io

import requests
from PIL import Image
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.shared import OxmlElement, qn

from utils import rebuild
from document_helpers import add_image_from_url, officialcastohead, officialcastobottom

def cert1chemichs(doc, arr):
    """Создает таблицу химического анализа для сварочных материалов"""
    table = doc.add_table(rows=9, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    tbl = table._tbl
    tbl_pr = tbl.tblPr


    tbl_borders = OxmlElement('w:tblBorders')
    border_types = ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']
    for border_type in border_types:
        border = OxmlElement(f'w:{border_type}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), '000000')
        tbl_borders.append(border)

    bottom_border = OxmlElement('w:bottom')
    bottom_border.set(qn('w:val'), 'single')
    bottom_border.set(qn('w:sz'), '4')
    bottom_border.set(qn('w:color'), '000000')
    tbl_borders.append(bottom_border)
    tbl_pr.append(tbl_borders)


    for col_idx in range(0, 6):
        if (col_idx % 2 == 0):
            bords = ['left']
        else:
            bords = ['right']
        for i, row in enumerate(table.rows):
            cell = row.cells[col_idx]
            tc_pr = cell._tc.get_or_add_tcPr()


            for elem in tc_pr.xpath(".//*[contains(name(), 'border')]"):
                tc_pr.remove(elem)

            if (i == 0):
                bords.append('top')
            elif (i == 8):
                bords.append('bottom')

            for border_type in bords:
                border = OxmlElement(f'w:{border_type}')
                border.set(qn('w:val'), 'single')
                border.set(qn('w:sz'), '18')
                border.set(qn('w:color'), '000000')
                tc_pr.append(border)
            bords = [bords[0]]


    analysis_data = [
        ("", "%", "", "%", "", "%"),
        ("С", arr.get("С", " "), "Cu", arr.get("Cu", " "), "Mo", arr.get("Mo", " ")),
        ("Si", arr.get("Si", " "), "V", arr.get("V", " "), "W", arr.get("W", " ")),
        ("Mn", arr.get("Mn", " "), "Fe", arr.get("Fe", " "), "В", arr.get("В", " ")),
        ("S", arr.get("S", " "), "Pb", arr.get("Pb", " "), "WC", arr.get("WC", " ")),
        ("P", arr.get("P", " "), "Sn", arr.get("Sn", " "), "Ti", arr.get("Ti", " ")),
        ("Cr", arr.get("Cr", " "), "Zn", arr.get("Zn", " "), "Mg", arr.get("Mg", " ")),
        ("Ni", arr.get("Ni", " "), "Al", arr.get("Al", " "), "N", arr.get("N", " ")),
        ("Nb", arr.get("Nb", " "), "Co", arr.get("Co", " "), "другие", arr.get("другие", " "))
    ]


    for row_idx in range(0, 9):
        row_cells = table.rows[row_idx].cells
        for col_idx in range(0, 6):
            row_cells[col_idx].text = analysis_data[row_idx][col_idx]
            if (col_idx % 2 == 1):
                row_cells[col_idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                row_cells[col_idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT


    for row in table.rows:
        tr = row._tr
        tr_height = OxmlElement('w:trHeight')
        tr_height.set(qn('w:val'), "270")
        tr_height.set(qn('w:hRule'), "exact")
        tr.append(tr_height)


    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)

    return table

def cert1table1(doc, arr):
    """Создает основную информационную таблицу для сварочных материалов"""
    table = doc.add_table(rows=7, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    tbl = table._tbl
    tbl_pr = tbl.tblPr


    tbl_borders = OxmlElement('w:tblBorders')
    border_types = ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']
    for border_type in border_types:
        border = OxmlElement(f'w:{border_type}')
        border.set(qn('w:val'), 'nil')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), '000000')
        tbl_borders.append(border)

    bottom_border = OxmlElement('w:bottom')
    bottom_border.set(qn('w:val'), 'nil')
    bottom_border.set(qn('w:sz'), '4')
    bottom_border.set(qn('w:color'), '000000')
    tbl_borders.append(bottom_border)
    tbl_pr.append(tbl_borders)


    analysis_data = [
        ('Тип продукта', 'Product type', arr.get("Тип продукта", " ")),
        ('Наименование', 'Designation', arr.get("Наименование", " ")),
        ('Диаметр, мм', 'Diameter, mm', arr.get("Диаметр, мм", " ")),
        ('Артикул', 'Article', str(arr.get("Артикул", " "))[:-2] if arr.get("Артикул") else " "),
        ('Номер партии', 'Batch No', arr.get("Номер партии", " ")),
        ('Количество, кг', 'Quantity, kg', arr.get("Количество, кг", " ")),
        ('Стандарт', 'Standard', arr.get("Стандарт", " "))
    ]


    for row_idx in range(0, 7):
        row_cells = table.rows[row_idx].cells
        for col_idx in range(3):
            row_cells[col_idx].text = analysis_data[row_idx][col_idx]


    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)


    for row in table.rows:
        tr = row._tr
        tr_height = OxmlElement('w:trHeight')
        tr_height.set(qn('w:val'), "470")
        tr_height.set(qn('w:hRule'), "exact")
        tr.append(tr_height)

    return table

def cert1table2(doc, arr):
    """Создает таблицу механических свойств для сварочных материалов"""
    table = doc.add_table(rows=5, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.width = Inches(7)

    tbl = table._tbl
    tbl_pr = tbl.tblPr


    tbl_layout = OxmlElement('w:tblLayout')
    tbl_layout.set(qn('w:type'), 'fixed')
    tbl_pr.append(tbl_layout)


    tbl_borders = OxmlElement('w:tblBorders')
    border_types = ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']
    for border_type in border_types:
        border = OxmlElement(f'w:{border_type}')
        border.set(qn('w:val'), 'nil')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), '000000')
        tbl_borders.append(border)

    bottom_border = OxmlElement('w:bottom')
    bottom_border.set(qn('w:val'), 'nil')
    bottom_border.set(qn('w:sz'), '4')
    bottom_border.set(qn('w:color'), '000000')
    tbl_borders.append(bottom_border)
    tbl_pr.append(tbl_borders)


    analysis_data = [
        ('Предел прочности', 'Tensile Strength (Rm):', arr.get("Предел прочности, МПа", " "),
         rebuild(arr.get("Предел прочности, МПа", " "), 'МПа')),
        ('Предел текучести', 'Yield Strength (Rp0.2):', arr.get("Предел текучести, МПа", " "),
         rebuild(arr.get("Предел текучести, МПа", " "), 'МПа')),
        ('Относительное удлинение', 'Elongation (A5):', arr.get("Относительное удлинение, %", " "),
         rebuild(arr.get("Относительное удлинение, %", " "), '%')),
        ('Ударная вязкость', 'Impact Strength:', arr.get("Ударная вязкость, Дж", " "),
         rebuild(arr.get("Ударная вязкость, Дж", " "), 'Дж')),
        ('Твердость', 'Hardness:', arr.get("Твердость", " "), ''),
    ]


    for row_idx in range(0, 5):
        row_cells = table.rows[row_idx].cells
        for col_idx in range(0, 4):
            row_cells[col_idx].text = analysis_data[row_idx][col_idx]
            if (col_idx == 2):
                row_cells[col_idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)


    for row in table.rows:
        tr = row._tr
        tr_height = OxmlElement('w:trHeight')
        tr_height.set(qn('w:val'), "500")
        tr_height.set(qn('w:hRule'), "exact")
        tr.append(tr_height)

    return table

def certificate1(arr):
    """Генерирует сертификат для сварочных материалов"""
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)


    add_image_from_url(
        doc,
        image_url="https://avatars.mds.yandex.net/i?id=93d97da040c3650f2e88feae2f0272b0d7071b44-5206685-images-thumbs&n=13",
        width=Inches(4)
    )


    officialcastohead(doc)
    doc.add_paragraph(" ", style='Normal')


    cert1table1(doc, arr)


    p2 = doc.add_paragraph()
    run2 = p2.add_run('Анализ (матрица) / Analysis (matrix)')
    run2.bold = True
    run2.font.color.rgb = RGBColor(0, 0, 0)


    cert1chemichs(doc, arr)


    cert1table2(doc, arr)


    try:
        neededlink = "https://disk.yandex.ru/i/RzqpHNoADQQLqA"
        file_id = "RzqpHNoADQQLqA"
        public_key = f"https://yadi.sk/i/{file_id}"
        api_url = "https://cloud-api.yandex.net/v1/disk/public/resources/download"
        response = requests.get(api_url, params={'public_key': public_key})

        if response.status_code == 200:
            download_data = response.json()
            download_url = download_data['href']
            image_response = requests.get(download_url)

            if image_response.status_code == 200:
                image_stream = io.BytesIO(image_response.content)
                img = Image.open(image_stream)
                width, height = img.size
                target_width = Inches(0.6)
                ratio = target_width / width
                target_height = height * ratio

                paragraph1 = doc.add_paragraph()
                paragraph1.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph1.add_run()
                run_text = paragraph1.add_run(' ' * 50 + 'Шмарова Л. О.' + ' ' * 50 + str(arr.get('Дата', ''))[:-9] if arr.get('Дата') else '')
                paragraph1.paragraph_format.tab_stops.add_tab_stop(Inches(3))
                run_text.bold = False
                run_text.font.color.rgb = RGBColor(0, 0, 0)
                run_text.font.size = Pt(10)
                run.add_picture(image_stream, width=target_width, height=target_height)
                paragraph1.paragraph_format.space_after = 0
    except:
        pass


    officialcastobottom(doc)


    doc.save('Сертификат1.docx')
    return 'Сертификат1.docx'

