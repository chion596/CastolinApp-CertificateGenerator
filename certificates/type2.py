import io

import requests
from PIL import Image
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.shared import OxmlElement, qn

from utils import rebuild
from document_helpers import add_image_from_url, officialcastohead, officialcastobottom

def cert23table1(doc, arr):
    """Создает основную информационную таблицу для порошков и керамики"""
    table = doc.add_table(rows=10, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    tbl = table._tbl
    tbl_pr = tbl.tblPr


    tbl_borders = OxmlElement('w:tblBorders')
    border_types = ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']
    for border_type in border_types:
        border = OxmlElement(f'w:{border_type}')
        border.set(qn('w:val'), 'nil')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), '000000')
        tbl_borders.append(border)

    bottom_border = OxmlElement('w:bottom')
    bottom_border.set(qn('w:val'), 'nil')
    bottom_border.set(qn('w:sz'), '4')
    bottom_border.set(qn('w:color'), '000000')
    tbl_borders.append(bottom_border)
    tbl_pr.append(tbl_borders)


    analysis_data = [
        ('Тип продукта', 'Product type', arr.get("Тип продукта", " ")),
        ('Наименование', 'Designation', arr.get("Наименование", " ")),
        ('Диаметр, мм', 'Diameter, mm', arr.get("Диаметр, мм", " ")),
        ('Артикул', 'Article', str(arr.get("Артикул", " "))[:-2] if arr.get("Артикул") else " "),
        ('Номер партии', 'Batch No', arr.get("Номер партии", " ")),
        ('Форма частиц', 'Grain form', arr.get("Форма частиц", " ")),
        ('Размер частиц', 'Grain sizerange', arr.get("Размер частиц", " ")),
        ('Текучесть', 'Hall flow', rebuild(arr.get("Текучесть", " "), arr.get("Текучесть", " ") + "    s/50g")),
        ('Насыпная плотность', 'Apparent density',
         rebuild(arr.get("Насыпная плотность", " "), arr.get("Насыпная плотность", " ") + "    g/cm3")),
        ('Твердость', 'Hardness', arr.get("Твердость", " "))
    ]


    for row_idx in range(0, 10):
        row_cells = table.rows[row_idx].cells
        for col_idx in range(3):
            row_cells[col_idx].text = analysis_data[row_idx][col_idx]


    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(12)


    for row in table.rows:
        tr = row._tr
        tr_height = OxmlElement('w:trHeight')
        tr_height.set(qn('w:val'), "485")
        tr_height.set(qn('w:hRule'), "exact")
        tr.append(tr_height)

    return table

def cert2chemichs(doc, arr):
    """Создает таблицу химического анализа для порошков"""
    table = doc.add_table(rows=8, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    tbl = table._tbl
    tbl_pr = tbl.tblPr


    tbl_borders = OxmlElement('w:tblBorders')
    border_types = ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']
    for border_type in border_types:
        border = OxmlElement(f'w:{border_type}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), '000000')
        tbl_borders.append(border)

    bottom_border = OxmlElement('w:bottom')
    bottom_border.set(qn('w:val'), 'single')
    bottom_border.set(qn('w:sz'), '4')
    bottom_border.set(qn('w:color'), '000000')
    tbl_borders.append(bottom_border)
    tbl_pr.append(tbl_borders)


    for col_idx in range(0, 6):
        if (col_idx % 2 == 0):
            bords = ['left']
        else:
            bords = ['right']
        for i, row in enumerate(table.rows):
            cell = row.cells[col_idx]
            tc_pr = cell._tc.get_or_add_tcPr()

            for elem in tc_pr.xpath(".//*[contains(name(), 'border')]"):
                tc_pr.remove(elem)

            if (i == 0):
                bords.append('top')
            elif (i == 7):
                bords.append('bottom')

            for border_type in bords:
                border = OxmlElement(f'w:{border_type}')
                border.set(qn('w:val'), 'single')
                border.set(qn('w:sz'), '18')
                border.set(qn('w:color'), '000000')
                tc_pr.append(border)
            bords = [bords[0]]


    analysis_data = [
        ("", "%", "", "%", "", "%"),
        ("С", arr.get("С", " "), "Ni", arr.get("Ni", " "), "WC", arr.get("WC", " ")),
        ("Cr", arr.get("Cr", " "), "Mo", arr.get("Mo", " "), "Ti", arr.get("Ti", " ")),
        ("Fe", arr.get("Fe", " "), "P", arr.get("P", " "), "Co", arr.get("Co", " ")),
        ("В", arr.get("В", " "), "Cu", arr.get("Cu", " "), "Al", arr.get("Al", " ")),
        ("Si", arr.get("Si", " "), "V", arr.get("V", " "), "Sn", arr.get("Sn", " ")),
        ("Zn", arr.get("Zn", " "), "Nb", arr.get("Nb", " "), "Pb", arr.get("Pb", " ")),
        ("Mn", arr.get("Mn", " "), "W", arr.get("W", " "), "другие", arr.get("другие", " ")),
    ]


    for row_idx in range(0, 8):
        row_cells = table.rows[row_idx].cells
        for col_idx in range(0, 6):
            row_cells[col_idx].text = analysis_data[row_idx][col_idx]
            if (col_idx % 2 == 1):
                row_cells[col_idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                row_cells[col_idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT


    for row in table.rows:
        tr = row._tr
        tr_height = OxmlElement('w:trHeight')
        tr_height.set(qn('w:val'), "325")
        tr_height.set(qn('w:hRule'), "exact")
        tr.append(tr_height)


    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(11)

    return table

def certificate2(arr):
    """Генерирует сертификат для порошков"""
    doc2 = Document()
    doc2.settings.odd_and_even_pages_header_footer = False


    add_image_from_url(
        doc2,
        image_url="https://avatars.mds.yandex.net/i?id=93d97da040c3650f2e88feae2f0272b0d7071b44-5206685-images-thumbs&n=13",
        width=Inches(4)
    )


    officialcastohead(doc2)
    doc2.add_paragraph(" ", style='Normal')


    cert23table1(doc2, arr)


    p = doc2.add_paragraph()
    run = p.add_run('Анализ (матрица) / Analysis (matrix)')
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)


    cert2chemichs(doc2, arr)
    doc2.add_paragraph(" ", style='Normal')


    try:
        neededlink = "https://disk.yandex.ru/i/RzqpHNoADQQLqA"
        file_id = "RzqpHNoADQQLqA"
        public_key = f"https://yadi.sk/i/{file_id}"
        api_url = "https://cloud-api.yandex.net/v1/disk/public/resources/download"
        response = requests.get(api_url, params={'public_key': public_key})

        if response.status_code == 200:
            download_data = response.json()
            download_url = download_data['href']
            image_response = requests.get(download_url)

            if image_response.status_code == 200:
                image_stream = io.BytesIO(image_response.content)
                img = Image.open(image_stream)
                width, height = img.size
                target_width = Inches(0.6)
                ratio = target_width / width
                target_height = height * ratio

                paragraph1 = doc2.add_paragraph()
                paragraph1.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph1.add_run()
                run_text = paragraph1.add_run(' ' * 50 + 'Шмарова Л. О.' + ' ' * 50 + str(arr.get('Дата', ''))[:-9] if arr.get('Дата') else '')
                paragraph1.paragraph_format.tab_stops.add_tab_stop(Inches(3))
                run_text.bold = False
                run_text.font.color.rgb = RGBColor(0, 0, 0)
                run_text.font.size = Pt(10)
                run.add_picture(image_stream, width=target_width, height=target_height)
                paragraph1.paragraph_format.space_after = 0
    except:
        pass


    officialcastobottom(doc2)


    doc2.save('Сертификат2.docx')
    return 'Сертификат2.docx'

