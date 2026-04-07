import io
from urllib.parse import urlparse

import requests
from PIL import Image
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.shared import OxmlElement, qn

def add_image_from_url(doc, image_url, width=Inches(3), caption=None):
    """Добавляет изображение из интернета в документ Word"""
    try:

        if not urlparse(image_url).scheme in ('http', 'https'):
            raise ValueError("Некорректный URL изображения")


        response = requests.get(image_url, timeout=10)
        response.raise_for_status()


        image_bytes = io.BytesIO(response.content)
        img = Image.open(image_bytes)


        width, height = img.size
        target_width = Inches(1.45)
        ratio = target_width / width
        target_height = height * ratio


        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run()
        run.add_picture(image_bytes, width=target_width, height=target_height)


        run_text = paragraph.add_run("\t" + 'Castolin Eutectic')
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(3))
        run_text.bold = True
        run_text.font.color.rgb = RGBColor(0, 0, 0)
        run_text.font.size = Pt(24)
        paragraph.paragraph_format.space_after = 0


        if caption:
            doc.add_paragraph(caption, style='Caption')

        return True
    except Exception as e:
        print(f"Ошибка загрузки изображения: {e}")
        return False

def officialcastohead(doc):
    """Создает верхний колонтитул сертификата"""
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = None
    table.autofit = False


    col_widths = [Inches(1.5), Inches(4), Inches(1.5)]
    for i, width in enumerate(col_widths):
        table.columns[i].width = width


    tbl = table._tbl
    tbl_pr = tbl.tblPr


    tbl_layout = OxmlElement('w:tblLayout')
    tbl_layout.set(qn('w:type'), 'fixed')
    tbl_pr.append(tbl_layout)


    tbl_borders = OxmlElement('w:tblBorders')


    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'single')
    top.set(qn('w:sz'), '32')
    top.set(qn('w:color'), '000000')
    tbl_borders.append(top)


    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '32')
    bottom.set(qn('w:color'), '000000')
    tbl_borders.append(bottom)


    for border_type in ['insideV']:
        border = OxmlElement(f'w:{border_type}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '32')
        border.set(qn('w:color'), '000000')
        tbl_borders.append(border)


    for border_type in ['left', 'right']:
        border2 = OxmlElement(f'w:{border_type}')
        border2.set(qn('w:val'), 'nil')
        tbl_borders.append(border2)

    tbl_pr.append(tbl_borders)


    for row in table.rows:
        tr = row._tr
        tr_height = OxmlElement('w:trHeight')
        tr_height.set(qn('w:val'), "750")
        tr_height.set(qn('w:hRule'), "exact")
        tr.append(tr_height)


    cells = table.rows[0].cells
    left_text = "EN 10 204"
    center_text = "ПРОТОКОЛ ИСПЫТАНИЙ\n\nTEST REPORT"
    right_text = "3.1"


    left_cell = cells[0]
    left_para = left_cell.paragraphs[0]
    left_run = left_para.add_run(left_text)
    left_run.bold = True
    left_run.font.size = Pt(18)
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER


    center_cell = cells[1]
    center_para = center_cell.paragraphs[0]
    center_run = center_para.add_run(center_text.split('\n')[0])
    center_run.bold = True
    center_run.font.size = Pt(11)

    center_para2 = center_cell.add_paragraph()
    center_run2 = center_para2.add_run(center_text.split('\n')[2])
    center_run2.bold = True
    center_run2.font.size = Pt(11)

    center_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    center_para2.alignment = WD_ALIGN_PARAGRAPH.CENTER


    right_cell = cells[2]
    right_para = right_cell.paragraphs[0]
    right_run = right_para.add_run(right_text)
    right_run.bold = True
    right_run.font.size = Pt(18)
    right_para.alignment = WD_ALIGN_PARAGRAPH.CENTER


    for cell in cells:
        cell.vertical_alignment = WD_ALIGN_PARAGRAPH.CENTER
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0

    return table

def officialcastobottom(doc):
    """Создает нижний колонтитул сертификата"""
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    table.autofit = False


    col_widths = [Inches(4.15), Inches(4.15)]
    for i, width in enumerate(col_widths):
        table.columns[i].width = width


    tbl = table._tbl
    tbl_pr = tbl.tblPr


    tbl_borders = OxmlElement('w:tblBorders')


    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'single')
    top.set(qn('w:sz'), '32')
    top.set(qn('w:color'), '000000')
    tbl_borders.append(top)


    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '32')
    bottom.set(qn('w:color'), '000000')
    tbl_borders.append(bottom)


    for border_type in ['insideV']:
        border = OxmlElement(f'w:{border_type}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '32')
        border.set(qn('w:color'), '000000')
        tbl_borders.append(border)


    for border_type in ['left', 'right']:
        border2 = OxmlElement(f'w:{border_type}')
        border2.set(qn('w:val'), 'nil')
        tbl_borders.append(border2)

    tbl_pr.append(tbl_borders)


    for row in table.rows:
        tr = row._tr
        tr_height = OxmlElement('w:trHeight')
        tr_height.set(qn('w:val'), "950")
        tr_height.set(qn('w:hRule'), "exactly")
        tr.append(tr_height)


    cells = table.rows[0].cells
    left_text = "ООО «Кастолин»"
    right_text = "ГАРАНТИЯ КАЧЕСТВА"


    left_cell = cells[0]
    left_para = left_cell.paragraphs[0]
    left_run = left_para.add_run(left_text)
    left_run.bold = True
    left_run.font.size = Pt(11)
    left_para.add_run().add_break()

    left_run2 = left_para.add_run('109147, г. Москва, ул. Таганская, д. 17-23')
    left_run2.font.size = Pt(11)
    left_para.add_run().add_break()

    left_run3 = left_para.add_run('Teл. +7 495 212 13 51')
    left_run3.font.size = Pt(11)
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER


    right_cell = cells[1]
    right_para = right_cell.paragraphs[0]
    right_run = right_para.add_run(right_text)
    right_run.bold = True
    right_run.font.size = Pt(12)
    right_para.add_run().add_break()

    right_run2 = right_para.add_run('QUALITY ASSURANCE')
    right_run2.bold = True
    right_run2.font.size = Pt(12)
    right_para.alignment = WD_ALIGN_PARAGRAPH.CENTER


    for cell in cells:
        cell.vertical_alignment = WD_ALIGN_PARAGRAPH.CENTER
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0

    return table

